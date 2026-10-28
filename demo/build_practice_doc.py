"""Regenerate demo/Mouse-AI-Practice.docx (needs: pip install python-docx).

The document is a short memo plus a numbered task list. Working through the
tasks in front of an audience produces a realistic burst of ribbon clicks,
which the overlay glows over and the dashboard tallies live.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Mouse-AI-Practice.docx"

TASKS = [
    "Select the title \"Quarterly Update\" and click Bold, then Italic.",
    "Change the title's font color using the Home tab.",
    "Highlight the word \"revenue\" in the first paragraph.",
    "Make the three bullet points below a numbered list.",
    "Insert a 3 x 3 table under \"Next steps\" (Insert > Table).",
    "Click Save (top-left) a couple of times while you work.",
    "Now move the pointer to the far corner of the screen and fast "
    "double-click: Mouse-AI returns it to your busiest spot (the ribbon).",
]


def build() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(12)

    banner = doc.add_paragraph()
    run = banner.add_run("Mouse-AI live demo: follow the steps, then check the dashboard.")
    run.italic = True
    run.font.color.rgb = RGBColor(0xC0, 0x40, 0x00)

    doc.add_heading("Tasks", level=2)
    for task in TASKS:
        doc.add_paragraph(task, style="List Number")

    doc.add_page_break()
    doc.add_paragraph("Quarterly Update", style="Title")
    doc.add_paragraph(
        "This quarter, revenue grew steadily across all regions. The support "
        "team closed more tickets than any previous quarter, and the new "
        "onboarding flow cut setup time roughly in half."
    )
    for item in ("Expand the pilot to two more teams",
                 "Review onboarding feedback",
                 "Plan the Q4 training sessions"):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Next steps", level=2)
    doc.add_paragraph("")

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
